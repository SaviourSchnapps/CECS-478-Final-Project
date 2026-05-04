"""Convert FINAL_REPORT.md to a DOCX.

Targeted converter — handles only the markdown subset used in the final
report (headings, paragraphs, fenced code, bullet/numbered lists, tables,
images, simple inline bold). Not a general-purpose markdown engine.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor


def _add_inline(paragraph, text: str) -> None:
    # very small inline pass: **bold**, `code`, links [text](url) -> text
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    pos = 0
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:  # `code`
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _flush_code(doc: Document, buf: list[str]) -> None:
    if not buf:
        return
    p = doc.add_paragraph()
    run = p.add_run("\n".join(buf))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.25)
    buf.clear()


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Light Grid Accent 1"
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            _add_inline(p, cell_text.strip())
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True


def convert(md_path: Path, docx_path: Path, image_root: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[list[str]] = []

    def _flush_table():
        nonlocal table_buf
        if table_buf:
            # drop the alignment row (---|---|---)
            table_buf = [r for r in table_buf if not all(set(c.strip()) <= set("-:| ") for c in r)]
            _add_table(doc, table_buf)
            table_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                _flush_code(doc, code_buf)
                in_code = False
            else:
                _flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if "|" in line and stripped.startswith("|") and stripped.endswith("|"):
            cells = [c for c in stripped.strip("|").split("|")]
            table_buf.append(cells)
            i += 1
            continue
        else:
            _flush_table()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph().add_run("─" * 60)
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2), level=min(level, 4))
            i += 1
            continue

        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)\s*$", stripped)
        if m:
            img_rel = m.group(1)
            img_path = (image_root / img_rel).resolve()
            if img_path.exists():
                doc.add_picture(str(img_path), width=Inches(5.0))
            else:
                p = doc.add_paragraph()
                _add_inline(p, f"[image not found: {img_rel}]")
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            doc.add_paragraph(style="List Number").add_run(m.group(2))
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, m.group(1))
            i += 1
            continue

        # paragraph: greedy join of subsequent non-blank lines
        para_lines = [stripped]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if not nxt or nxt.startswith(("#", "```", "- ", "* ", "!", "|"))\
                    or re.match(r"^\d+\.\s", nxt):
                break
            para_lines.append(nxt)
            j += 1
        p = doc.add_paragraph()
        _add_inline(p, " ".join(para_lines))
        i = j

    _flush_code(doc, code_buf)
    _flush_table()

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"wrote {docx_path}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--in", dest="src", type=Path, required=True)
    ap.add_argument("--out", dest="dst", type=Path, required=True)
    ap.add_argument("--image-root", type=Path, default=None,
                    help="directory relative paths in ![](...) are resolved against. "
                         "Defaults to the source markdown's parent directory.")
    args = ap.parse_args()
    image_root = args.image_root or args.src.parent
    convert(args.src, args.dst, image_root)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
